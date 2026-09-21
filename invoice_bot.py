# =============================================================================
# invoice_bot.py — Main Invoice Processing Engine
# Invoice Management System
# Version 2.0 — With Backup, Database, and Duplicate Detection
# =============================================================================

import os
import xmltodict
import re
from datetime import datetime, timedelta
import openpyxl
from copy import copy
from openpyxl.styles import Alignment
import shutil

import config
import database
import backup


# ─── HELPER FUNCTIONS ─────────────────────────────────────────────────────────

def calculate_business_day(issue_date: datetime) -> datetime:
    """Calculates due date in 7 days, skipping weekends."""
    due_date = issue_date + timedelta(days=7)
    if due_date.weekday() == 5:   # Saturday → Monday
        return due_date + timedelta(days=2)
    elif due_date.weekday() == 6:  # Sunday → Monday
        return due_date + timedelta(days=1)
    return due_date


def apply_custom_style(source_cell, dest_cell, col: str):
    """Copies formatting style from reference cell to destination cell (Optimized Version)."""
    if source_cell.has_style:
        # Direct and ultra-fast base style assignment
        dest_cell._style = source_cell._style

        # Overwrites only specific formatting if necessary
        if col == "J":
            dest_cell.alignment     = Alignment(horizontal='right', vertical='bottom')
            dest_cell.number_format = '_-* #,##0.00_-;_-* #,##0.00_-;_-* "-"??_-;_-@_-'
        elif col == "T":
            dest_cell.alignment     = Alignment(vertical='bottom')
            dest_cell.number_format = 'dd/mm/yyyy'
        elif col == "F":
            dest_cell.number_format = 'dd/mm/yy'
            dest_cell.alignment     = Alignment(horizontal='right', vertical='bottom')
        elif col == "Z":
            dest_cell.alignment     = Alignment(horizontal='center', vertical='bottom')


def _extract_xml_data(info: dict, item: dict, invoice_num: str, cfg: dict) -> dict:
    """Extracts and parses data from an invoice item (Turbo Version)."""
    xProd    = str(item['prod']['xProd']).upper()
    issue_dt = datetime.strptime(info['ide']['dhEmi'][:10], "%Y-%m-%d")
    totals   = info['total']['ICMSTot']

    # Acronym and Quality Logic
    #
    # The keywords cover both the original Portuguese descriptions and the
    # English ones used in the anonymized sample files, so the same rule
    # classifies either. BZ = galvanized, BF = cold rolled, BQ = hot rolled.
    if any(x in xProd for x in ["GALV", "ZC", "NBR", "7008"]):
        acronym = "BZ"
        m = re.search(r'(NBR\s?\d+[^,]*|ZC\s?\d+)', xProd)
        quality = m.group(1) if m else "NBR 7008 ZC"
    elif any(x in xProd for x in ["FRIO", "COLD ROLLED", "BF"]):
        acronym = "BF"
        m = re.search(r'(SAE\s?J?\d+\s?\d*)', xProd)
        quality = m.group(1) if m else "SAE 1008"
    else:
        acronym = "BQ"
        m = re.search(r'(SAE\s?J?\d+\s?\d*|A36)', xProd)
        quality = m.group(1) if m else "SAE"

    m_thick = re.search(r'(\d+\.\d{3})', xProd)
    thickness = float(m_thick.group(1)) if m_thick else 0.0

    m_width = re.search(r'X\s?(\d+)\.', xProd)
    width   = int(m_width.group(1)) if m_width else ""

    m_order = re.search(r'(\d{10})', xProd)
    order_num = m_order.group(1) if m_order else ""

    return {
        "INVOICE"    : int(invoice_num),
        "ORDER"      : order_num,
        "SUPPLIER"   : cfg.get("default_supplier", 36003),
        "INV_DATE"   : issue_dt,
        "WEIGHT"     : int(float(item['prod']['qCom']) * 1000),
        "PROD"       : acronym,
        "QUALITY"    : quality,
        "THICKNESS"  : thickness,
        "WIDTH"      : width,
        "TOTAL_VALUE": float(totals['vNF']),
        "TAX_ICMS"   : float(totals['vICMS']),
        "TAX_IPI"    : float(totals['vIPI']),
        "DUE_DATE"   : calculate_business_day(issue_dt),
        "BILLING"    : f"00{invoice_num}01",
        "INV_VALUE"  : float(totals['vNF']),
        "NCM"        : item['prod'].get('NCM', '')
    }


def _extract_access_key(xml_data: dict) -> str:
    """Extracts the invoice access key for duplicate control."""
    try:
        nfe_info = xml_data.get('nfeProc', {}).get('NFe', {}).get('infNFe', {})
        access_key = nfe_info.get('@Id', '').replace('NFe', '')
        if not access_key:
            # Fallback: uses the protocol key
            access_key = xml_data.get('nfeProc', {}).get('protNFe', {}).get(
                'infProt', {}).get('chNFe', '')
        return access_key
    except Exception:
        return ""


# ─── MAIN FUNCTION ────────────────────────────────────────────────────────────

def process_files(callback=None):
    """
    Main engine: reads XMLs, matches with spreadsheet by batch,
    saves data preserving formatting, backups before saving,
    and logs everything in the database.
    """
    cfg = config.load()
    xml_folder       = cfg.get("xml_folder", "./input_xml")
    processed_folder = cfg.get("processed_folder", "./processed_xml")
    excel_file       = cfg.get("excel_file", "./production_planning.xlsx")
    batch_col        = cfg.get("batch_column", "W")
    sheet_name       = cfg.get("sheet_name", "Steel Plant")
    column_map       = cfg.get("column_map", 
    {"INVOICE": "C", 
    "ORDER": "D", 
    "SUPPLIER": "E", 
    "INV_DATE": "F",
    "WEIGHT": "G", 
    "PROD": "H", 
    "QUALITY": "I", 
    "THICKNESS": "J", 
    "WIDTH": "K",
    "TOTAL_VALUE": "L",
    "TAX_ICMS": "M", 
    "TAX_IPI": "N",
    "DUE_DATE": "T", 
    "BILLING": "U", 
    "INV_VALUE": "V",
    "NCM": "Z"})

    database.initialize()

    def _log(msg: str):
        if callback:
            callback(msg)

    # ── Initial Validations ──
    if not excel_file or not os.path.exists(excel_file):
        _log("⚠️ Excel spreadsheet not found! Check path in Settings.")
        database.log("ERROR", "Excel spreadsheet not found.", excel_file)
        return

    if not xml_folder or not os.path.exists(xml_folder):
        _log("⚠️ Pending XMLs folder not found!")
        database.log("ERROR", "XML folder not found.", xml_folder)
        return

    # ── Open Spreadsheet ──
    try:
        _log("⏳ Opening spreadsheet... please wait.")
        wb    = openpyxl.load_workbook(excel_file, data_only=False)
        sheet = next((wb[n] for n in wb.sheetnames if sheet_name in n), None)

        if sheet is None:
            _log(f"⚠️ Tab '{sheet_name}' not found in spreadsheet!")
            database.log("ERROR", f"Tab '{sheet_name}' not found.")
            return

    except Exception as e:
        _log(f"⚠️ Error opening Excel: {e}")
        database.log("ERROR", f"Error opening Excel: {e}")
        return

    # ── List XMLs ──
    xml_files = [f for f in os.listdir(xml_folder) if f.lower().endswith(".xml")]
    if not xml_files:
        _log("ℹ️ No XML files in pending folder.")
        return

    _log(f"📂 {len(xml_files)} XML(s) found. Starting processing...")

# ── TURBO: In-memory batch map (OPTIMIZED) ──
    excel_batch_map = {}
    # Gets column letter (e.g., 'W') and iterates only to the last row with data
    max_row = sheet.max_row
    
    for row in range(1, max_row + 1):
        cell = sheet[f"{batch_col}{row}"]
        if cell.value:
            val = str(cell.value).strip().upper()
            excel_batch_map[val] = row

    files_to_move    = []
    was_updated      = False
    backup_done      = False
    total_processed  = 0
    total_duplicates = 0
    total_not_found  = 0

    for file_name in xml_files:
        xml_path = os.path.join(xml_folder, file_name)

        try:
            with open(xml_path, "rb") as f:
                xml_data = xmltodict.parse(f.read())

            info        = xml_data['nfeProc']['NFe']['infNFe']
            invoice_num = info['ide']['nNF']
            access_key  = _extract_access_key(xml_data)

            # ── PROTECTION: duplicate detection ──
            if access_key and database.invoice_already_processed(access_key):
                _log(f"⚠️  Invoice {invoice_num} was already processed — skipping.")
                database.log("WARNING", f"Duplicate invoice skipped: {invoice_num}", access_key)
                total_duplicates += 1
                # Moves to processed folder anyway to clean up pending folder
                if processed_folder:
                    shutil.move(xml_path, os.path.join(processed_folder, file_name))
                continue

            details = info['det']
            if not isinstance(details, list):
                details = [details]

            # ── Extract Batch ──
            xml_batch = ""
            try:
                transp    = info.get('transp', {})
                vol       = transp.get('vol', {})
                xml_batch = str(
                    vol[0].get('nVol', '') if isinstance(vol, list) else vol.get('nVol', '')
                ).upper().strip()
            except Exception:
                xml_batch = ""

            found_row = excel_batch_map.get(xml_batch)

            if found_row:
                # ── PROTECTION: creates backup ONCE before first save ──
                if not backup_done:
                    _log("🔒 Backing up spreadsheet before making changes...")
                    bkp_result = backup.create_backup(reason=f"Auto before logging Inv {invoice_num}")
                    if bkp_result:
                        _log(f"✅ Backup created: {os.path.basename(bkp_result)}")
                    else:
                        _log("⚠️  Could not create backup. Continuing anyway.")
                    backup_done = True

                total_invoice_value = 0.0

                # ── Process only the first item to avoid slow rewriting ──
                for item in details:
                    # Pass cfg loaded at top to avoid disk I/O
                    data = _extract_xml_data(info, item, invoice_num, cfg)
                    total_invoice_value = data["TOTAL_VALUE"]

                    for ref, col in column_map.items():
                        dest_cell  = sheet[f"{col}{found_row}"]
                        dest_cell.value = data.get(ref)
                        if found_row > 1:
                            apply_custom_style(
                                sheet[f"{col}{found_row - 1}"],
                                dest_cell,
                                col
                            )
                    
                    # Break after 1st item. Since batch points to a SINGLE row in Excel, 
                    # processing other items of same invoice would just freeze PC rewriting same row.
                    break

                # ── Log in database ──
                database.register_invoice(
                    invoice_num = invoice_num,
                    access_key  = access_key,
                    supplier    = str(data.get("SUPPLIER", "")),
                    total_value = total_invoice_value,
                    inv_date    = data.get("INV_DATE", datetime.now()).strftime("%Y-%m-%d"),
                    batch       = xml_batch,
                    excel_row   = found_row,
                    xml_file    = file_name
                )
                database.log("SUCCESS", f"Invoice {invoice_num} processed on row {found_row}.", xml_batch)

                _log(f"🚀 Invoice {invoice_num} processed → Row {found_row} (Batch: {xml_batch})")
                was_updated = True
                total_processed += 1
                files_to_move.append(file_name)

            else:
                _log(f"❌ Invoice {invoice_num}: Batch '{xml_batch}' not found in spreadsheet.")
                database.log("WARNING", f"Batch not found: Inv {invoice_num}", xml_batch)
                total_not_found += 1

        except Exception as e:
            _log(f"⚠️  Error in file {file_name}: {e}")
            database.log("ERROR", f"Error processing XML: {file_name}", str(e))

    # ── TURBO: Single final save ──
    if was_updated:
        try:
            _log("💾 Saving all invoices to Excel... please wait.")
            wb.save(excel_file)
            database.log("INFO", f"{total_processed} Invoice(s) saved to Excel.")

            # Move XMLs only after successful save
            for file_name in files_to_move:
                if processed_folder:
                    shutil.move(
                        os.path.join(xml_folder, file_name),
                        os.path.join(processed_folder, file_name)
                    )

            _log(f"\n{'─'*45}")
            _log(f"🏁 PROCESS COMPLETED!")
            _log(f"   ✅ Processed : {total_processed}")
            if total_duplicates:
                _log(f"   ⚠️  Duplicates: {total_duplicates}")
            if total_not_found:
                _log(f"   ❌ Not found : {total_not_found}")
            _log(f"{'─'*45}\n")

        except Exception as e:
            _log(f"❌ ERROR SAVING EXCEL: {e}")
            database.log("ERROR", f"Error saving Excel: {e}")
    else:
        _log("ℹ️ Process finished. No changes made to the spreadsheet.")


if __name__ == "__main__":
    process_files(callback=print)